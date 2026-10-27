"""Parte de nómina sobre el archivo oficial, sin alterar su formato.

El archivo enviado es la plantilla original de Excel con solo los datos de la semana
modificados directamente en su XML: encabezado de la semana, números de día, asistencia
y los valores calculados de las fórmulas existentes. Tabla, validaciones, estilos,
fórmulas y configuración quedan idénticos, y el archivo trae los resultados ya
calculados para que cualquier sistema que lo importe lea los mismos valores que Excel.
"""
import math
import re
import shutil
import zipfile
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

from . import semanas

PLANTILLA = Path(__file__).resolve().parent.parent / "plantillas" / "nomina_plantilla.xlsx"

COLS_DIAS = ["E", "G", "I", "K", "M"]          # lunes a viernes
COLS_SEMANA = ["E", "G", "I", "K", "M", "O", "Q"]  # números de día, lunes a domingo
VALORES_PERMITIDOS = (0.0, 0.5, 1.0)  # validación de datos del formato oficial

PERSONAL = {
    "german": {"fila": 17, "nombre": "Germán", "jornada_horas": 8,
               "nombre_completo": "Germán Rojas Galeana", "puesto": "Lead Software Architect"},
    "aldahir": {"fila": 18, "nombre": "Aldahir", "jornada_horas": 4,
                "nombre_completo": "Aldahir Molina Medina", "puesto": "Junior Software Developer"},
}

MESES_MAY = [m.upper() for m in semanas.MESES]


# ---------- reglas de asistencia ----------

def horas_a_asistencia(horas: list[float], jornada: float) -> list[float]:
    """Convierte horas a 0, 0.5 o 1, redondeando hacia abajo al medio día."""
    return [min(math.floor(max(h, 0) / jornada * 2 + 1e-9) / 2, 1.0) for h in horas]


def salario_semanal(clave: str) -> float:
    ws = openpyxl.load_workbook(PLANTILLA, read_only=True).worksheets[0]
    return float(ws[f"V{PERSONAL[clave]['fila']}"].value or 0)


def _redondear(v: float, dec: int = 2) -> float:
    return float(Decimal(str(v)).quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP))


def _num(v: float) -> str:
    return f"{v:.15g}"


def encabezado_semana(lunes: date) -> str:
    """Mismo patrón del original: 'SEMANA DEL  19 AL 23  DE  OCTUBRE DEL 2026'."""
    v = semanas.viernes(lunes)
    if lunes.month == v.month:
        return f"SEMANA DEL  {lunes.day} AL {v.day}  DE  {MESES_MAY[v.month - 1]} DEL {v.year}"
    if lunes.year == v.year:
        return (f"SEMANA DEL  {lunes.day} DE {MESES_MAY[lunes.month - 1]} AL {v.day}  DE  "
                f"{MESES_MAY[v.month - 1]} DEL {v.year}")
    return (f"SEMANA DEL  {lunes.day} DE {MESES_MAY[lunes.month - 1]} DEL {lunes.year} AL "
            f"{v.day}  DE  {MESES_MAY[v.month - 1]} DEL {v.year}")


def calcular(asistencia: dict[str, list[float]]) -> dict[str, float]:
    """Replica las fórmulas de la plantilla (filas 17 y 18 y totales de la fila 22)."""
    c: dict[str, float] = {}
    for clave, valores in asistencia.items():
        f = PERSONAL[clave]["fila"]
        s = sum(valores)
        v = salario_semanal(clave)
        w = _redondear((v / 5) * s)
        x = 0.0  # sin tiempo extra capturado (columnas EX vacías)
        c[f"S{f}"], c[f"T{f}"], c[f"U{f}"] = s, s * 8, 0.0
        c[f"W{f}"], c[f"X{f}"] = w, x
        c[f"AA{f}"] = _redondear(w + x)
    c["W22"] = _redondear(c["W17"] + c["W18"])
    c["X22"] = 0.0
    c["AA22"] = _redondear(c["AA17"] + c["AA18"])
    return c


# ---------- edición directa del XML ----------

def _poner_valor(xml: str, ref: str, valor: str) -> str:
    patron = re.compile(rf'<c r="{ref}"(?P<attrs>[^>]*?)(?:/>|>(?P<cuerpo>.*?)</c>)', re.S)
    m = patron.search(xml)
    if not m:
        raise ValueError(f"No existe la celda {ref} en la plantilla")
    attrs, cuerpo = m.group("attrs"), m.group("cuerpo")
    if cuerpo is None:
        nuevo = f'<c r="{ref}"{attrs}><v>{valor}</v></c>'
    elif re.search(r"<v>.*?</v>|<v/>", cuerpo, re.S):
        nuevo = f'<c r="{ref}"{attrs}>' + re.sub(r"<v>.*?</v>|<v/>", f"<v>{valor}</v>", cuerpo, 1, re.S) + "</c>"
    else:
        nuevo = f'<c r="{ref}"{attrs}>{cuerpo}<v>{valor}</v></c>'
    return xml[:m.start()] + nuevo + xml[m.end():]


def _indice_cadena(sheet_xml: str, ref: str) -> int:
    m = re.search(rf'<c r="{ref}"[^>]*t="s"[^>]*><v>(\d+)</v></c>', sheet_xml)
    if not m:
        raise ValueError(f"{ref} no es texto compartido en la plantilla")
    return int(m.group(1))


def _reemplazar_cadena(ss_xml: str, indice: int, texto: str) -> str:
    items = list(re.finditer(r"<si>.*?</si>", ss_xml, re.S))
    m = items[indice]
    original = m.group(0)
    t = re.fullmatch(r"<si>(<t[^>]*>).*?</t></si>", original, re.S)
    apertura = t.group(1) if t else "<t>"
    if texto != texto.strip() and "xml:space" not in apertura:
        apertura = '<t xml:space="preserve">'
    return ss_xml[:m.start()] + f"<si>{apertura}{escape(texto)}</t></si>" + ss_xml[m.end():]


def generar_archivo(lunes: date, asistencia: dict[str, list[float]], destino: Path) -> Path:
    titulo = semanas.nombre_hoja(lunes)
    valores = calcular(asistencia)
    with zipfile.ZipFile(PLANTILLA) as zin:
        hoja_actual = re.search(r'<sheet name="([^"]+)"', zin.read("xl/workbook.xml").decode()).group(1)
        contenido = {i.filename: zin.read(i.filename) for i in zin.infolist()}
        infos = zin.infolist()

    sheet = contenido["xl/worksheets/sheet1.xml"].decode("utf-8")
    ss = contenido["xl/sharedStrings.xml"].decode("utf-8")

    ss = _reemplazar_cadena(ss, _indice_cadena(sheet, "B11"), encabezado_semana(lunes))
    for i, col in enumerate(COLS_SEMANA):
        sheet = _poner_valor(sheet, f"{col}16", str((lunes + timedelta(days=i)).day))
    for clave, dias in asistencia.items():
        fila = PERSONAL[clave]["fila"]
        for col, v in zip(COLS_DIAS, dias):
            sheet = _poner_valor(sheet, f"{col}{fila}", _num(v))
    for ref, v in valores.items():
        sheet = _poner_valor(sheet, ref, _num(v))

    contenido["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    contenido["xl/sharedStrings.xml"] = ss.encode("utf-8")
    nombre_esc = escape(hoja_actual)
    for parte in ("xl/workbook.xml", "docProps/app.xml"):
        if parte in contenido:
            contenido[parte] = contenido[parte].decode("utf-8").replace(
                nombre_esc, escape(titulo)).encode("utf-8")

    destino.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destino, "w") as zout:
        for info in infos:
            zout.writestr(info, contenido[info.filename], compress_type=info.compress_type)
    return destino


def generar_semana(lunes: date, asistencia: dict[str, list[float]], data_dir: Path) -> Path:
    """Crea el archivo de la semana (el que se envía) y actualiza el maestro interno."""
    for clave, dias in asistencia.items():
        if any(v not in VALORES_PERMITIDOS for v in dias):
            raise ValueError(f"Asistencia de {clave} fuera de 0, 0.5 o 1: {dias}")
    archivo = generar_archivo(lunes, asistencia, data_dir / "semanas" / semanas.nombre_archivo(lunes))
    _actualizar_maestro(lunes, asistencia, data_dir / "nomina_maestro.xlsx")
    return archivo


# ---------- maestro interno (historial para consulta, no se envía) ----------

def _llenar_maestro(ws, lunes: date, asistencia: dict[str, list[float]]) -> None:
    ws["B11"] = encabezado_semana(lunes)
    for i, col in enumerate(COLS_SEMANA):
        ws[f"{col}16"] = (lunes + timedelta(days=i)).day
    for clave, valores in asistencia.items():
        fila = PERSONAL[clave]["fila"]
        for col, valor in zip(COLS_DIAS, valores):
            ws[f"{col}{fila}"] = valor
    ws["AC16"] = datetime(lunes.year, lunes.month, lunes.day)  # clave de orden, fuera del área impresa
    ws.column_dimensions["AC"].hidden = True


def _actualizar_maestro(lunes: date, asistencia: dict[str, list[float]], maestro: Path) -> None:
    if not maestro.exists():
        shutil.copy(PLANTILLA, maestro)
        wb = openpyxl.load_workbook(maestro)
        wb.worksheets[0].title = "PLANTILLA"
    else:
        wb = openpyxl.load_workbook(maestro)
    titulo = semanas.nombre_hoja(lunes)
    if titulo in wb.sheetnames:
        del wb[titulo]
    ws = wb.copy_worksheet(wb["PLANTILLA"])
    ws.title = titulo
    _llenar_maestro(ws, lunes, asistencia)
    hojas = sorted((w for w in wb.worksheets if w.title != "PLANTILLA"), key=lambda w: w["AC16"].value)
    wb._sheets = [wb["PLANTILLA"], *hojas]
    wb["PLANTILLA"].sheet_state = "hidden"
    wb.active = wb._sheets.index(ws)
    for w in wb.worksheets:
        w.sheet_view.tabSelected = w is ws
    wb.calculation.fullCalcOnLoad = True
    wb.save(maestro)


def resumen(asistencia: dict[str, list[float]]) -> tuple[str, float]:
    """Texto de vista previa y total, idénticos a lo que calcula el archivo."""
    c = calcular(asistencia)
    lineas = []
    dias = ["L", "M", "Mi", "J", "V"]
    for clave, valores in asistencia.items():
        p = PERSONAL[clave]
        f = p["fila"]
        detalle = "  ".join(f"{d} {v:g}" for d, v in zip(dias, valores))
        lineas.append(f"{p['nombre']}: {detalle}\n   {c[f'S{f}']:g} días · ${c[f'AA{f}']:,.2f}")
    return "\n".join(lineas), c["AA22"]
